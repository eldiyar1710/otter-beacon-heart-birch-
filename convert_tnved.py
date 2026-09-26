#!/usr/bin/env python3
"""
Скрипт для конвертации Excel-файла ТН ВЭД (с tws.by) в JSON для калькулятора.

Использование:
1. Скачайте файл: https://www.tws.by/tws/tnved/download/excel
2. Положите его рядом со скриптом как tnved.xlsx
3. Запустите: python convert_tnved.py
4. Получите tnved_data.js — подключите его к калькулятору

Требуется: pip install openpyxl
"""

import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Установите openpyxl: pip install openpyxl")
    sys.exit(1)

EXCEL_FILE = Path("tnved.xlsx")
OUTPUT_JS = Path("tnved_data.js")
OUTPUT_JSON = Path("tnved_data.json")

def detect_columns(ws):
    headers = {}
    for col in range(1, min(ws.max_column + 1, 20)):
        val = ws.cell(1, col).value
        if val is None:
            continue
        val_lower = str(val).lower().strip()
        if "код" in val_lower or "code" in val_lower or "tnved" in val_lower:
            headers["code"] = col
        elif "наимен" in val_lower or "name" in val_lower or "описан" in val_lower or "товар" in val_lower:
            headers["name"] = col
        elif "ставк" in val_lower or "пошлин" in val_lower or "тариф" in val_lower or "rate" in val_lower:
            headers["rate"] = col
    return headers

def parse_rate(rate_str):
    if rate_str is None:
        return {"type": "advalorem", "value": 0, "text": "0%"}
    s = str(rate_str).strip().replace(",", ".")
    if not s or s in ("-", "—", "нет"):
        return {"type": "advalorem", "value": 0, "text": "0%"}
    if "%" in s and "евро" not in s.lower() and "eur" not in s.lower() and "не менее" not in s.lower():
        try:
            num = float(s.replace("%", "").strip())
            return {"type": "advalorem", "value": num, "text": f"{num}%"}
        except:
            pass
    if "не менее" in s.lower():
        try:
            import re
            pct = re.search(r"(\d+[.,]?\d*)\s*%", s)
            specific = re.search(r"(\d+[.,]?\d*)\s*(евро|eur|€)", s.lower())
            pct_val = float(pct.group(1).replace(",", ".")) if pct else 15
            spec_val = float(specific.group(1).replace(",", ".")) if specific else 0
            return {"type": "combined", "value": pct_val, "specific": spec_val, "text": s}
        except:
            pass
    return {"type": "advalorem", "value": 0, "text": s}

def convert():
    if not EXCEL_FILE.exists():
        print(f"Файл {EXCEL_FILE} не найден!")
        print("Скачайте: https://www.tws.by/tws/tnved/download/excel")
        return False
    print(f"Открываю {EXCEL_FILE}...")
    wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True, data_only=True)
    ws = wb.active
    headers = detect_columns(ws)
    if "code" not in headers:
        headers = {"code": 1, "name": 2, "rate": 3}
    data = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        code_cell = row[headers["code"] - 1].value
        if code_cell is None:
            continue
        code = str(code_cell).replace(" ", "").replace(".", "").strip()
        if not code.isdigit() or len(code) < 4:
            continue
        name = ""
        if "name" in headers:
            name_val = row[headers["name"] - 1].value
            name = str(name_val).strip() if name_val else ""
        rate_raw = row[headers["rate"] - 1].value if "rate" in headers else None
        rate_info = parse_rate(rate_raw)
        vat = 10 if code.startswith(("02", "03", "04", "07", "08", "09", "10", "11", "12", "15", "16", "17", "18", "19", "20", "21")) else 20
        item = {
            "code": code.ljust(10, "0")[:10] if len(code) <= 10 else code[:10],
            "name": name[:200],
            "rate": rate_info["text"],
            "rateType": rate_info["type"],
            "rateValue": rate_info.get("value", 0),
            "specific": rate_info.get("specific"),
            "vat": vat,
            "unit": "кг",
            "cert": "декларация" if code.startswith(("16", "17", "18", "19", "20", "21")) else "уточнять",
            "certInfo": "Проверьте по ТР ТС на tnved.info"
        }
        data.append(item)
    wb.close()
    print(f"Загружено кодов: {len(data)}")
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    with open(OUTPUT_JS, "w", encoding="utf-8") as f:
        f.write("const TNVED_FULL_DB = ")
        json.dump(data, f, ensure_ascii=False)
        f.write(";\n")
    print(f"Сохранены {OUTPUT_JSON} и {OUTPUT_JS}")
    return True

if __name__ == "__main__":
    convert()
